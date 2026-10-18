import csv
import json
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook
from selenium.common.exceptions import (
    ElementClickInterceptedException,
    StaleElementReferenceException,
    TimeoutException,
)
from selenium.webdriver.common.keys import Keys

from automacao import (
    DEFAULT_WHATSAPP_GROUP,
    LEGACY_REPORT_FIELDS,
    REPORT_FIELDS,
    load_note_metadata,
    load_notes,
    normalize_note,
    parse_args,
    prepare_report,
    select_notes_to_process,
    write_note_descriptions,
    write_report,
)
from salesforce import (
    ResultadoNota,
    _abrir_compromisso_servico,
    _abrir_aba_relacionados,
    _baixar_arquivos_relacionados,
    _compromissos_relacionados,
    _extract_downloaded_images,
    _matching_work_order_result,
    _modal_image_rows,
    _modal_download_button,
    _processar_nota,
    _visible_search_box,
    _wait_for_stable_search_box,
    _buscar_ordem,
    consultar_notas,
    criar_navegador,
)
from whatsapp import (
    _note_message,
    _find_file_input,
    _click_attachment_button,
    _click_media_menu_item,
    _media_dialog,
    _open_group,
    _wait_for_document_sent,
    _send_button,
    _send_media,
    _send_images_with_caption,
    _xpath_literal,
    create_responses_archive,
    enviar_respostas,
)


class NormalizeNoteTests(unittest.TestCase):
    def test_normalizes_integer_and_integral_float(self):
        self.assertEqual(normalize_note(12345, 2), "12345")
        self.assertEqual(normalize_note(12345.0, 2), "12345")

    def test_trims_text_and_ignores_blank_values(self):
        self.assertEqual(normalize_note(" 00123 ", 2), "00123")
        self.assertIsNone(normalize_note(" ", 2))
        self.assertIsNone(normalize_note(None, 2))

    def test_rejects_invalid_numbers(self):
        with self.assertRaises(ValueError):
            normalize_note(12.5, 2)
        with self.assertRaises(ValueError):
            normalize_note(True, 2)
        with self.assertRaisesRegex(ValueError, "somente algarismos"):
            normalize_note("../fora", 2)


class WorkbookTests(unittest.TestCase):
    def test_loads_column_case_insensitively_and_deduplicates(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notas.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Notas sem evidências"
            sheet.append(["NOTA", "OBS"])
            sheet.append([123, "a"])
            sheet.append([123.0, "duplicada"])
            sheet.append([" 00456 ", "b"])
            sheet.append([None, "vazia"])
            workbook.save(path)

            self.assertEqual(
                load_notes(path, "Notas sem evidências", "Nota"),
                ["123", "00456"],
            )

    def test_reports_missing_sheet_or_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notas.xlsx"
            workbook = Workbook()
            workbook.active.append(["Outra coluna"])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "Aba"):
                load_notes(path, "Inexistente", "Nota")
            with self.assertRaisesRegex(ValueError, "Coluna"):
                load_notes(path, "Sheet", "Nota")

    def test_loads_metadata_and_writes_one_description_per_note(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "notas.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Notas sem evidências"
            sheet.append(
                [
                    "Nota",
                    "Texto code ação",
                    "Código de ação",
                    "Data encermto.",
                    "CenTrab respon.",
                    "Texto code ação",
                    "UTD",
                ]
            )
            sheet.append(
                [
                    123,
                    "Não Localizado",
                    "NE21",
                    datetime(2026, 9, 28),
                    "4DW501",
                    "Não Localizado",
                    "UTD-01",
                ]
            )
            workbook.save(path)

            metadata = load_note_metadata(
                path,
                "Notas sem evidências",
                "Nota",
                ["123"],
            )
            output = root / "Respostas"
            (output / "123").mkdir(parents=True)
            write_note_descriptions(output, ["123"], metadata)

            self.assertEqual(metadata["123"]["data_encerramento"], "28/09/2026")
            self.assertEqual(metadata["123"]["UTD"], "UTD-01")
            self.assertEqual(
                (output / "123" / "descricao.txt").read_text(encoding="utf-8"),
                "Nota: 123\n"
                "Texto code ação: Não Localizado\n"
                "Código de ação: NE21\n"
                "Data de encerramento: 28/09/2026\n"
                "Centro de trabalho responsável: 4DW501\n",
            )

    def test_rejects_conflicting_duplicate_description_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notas.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Notas sem evidências"
            sheet.append(
                [
                    "Nota",
                    "Texto code ação",
                    "Código de ação",
                    "Data encermto.",
                    "CenTrab respon.",
                    "Texto code ação",
                    "UTD",
                ]
            )
            sheet.append([123, "A", "NE01", None, "4DW501", "B", "UTD-01"])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "duplicadas"):
                load_note_metadata(path, "Notas sem evidências", "Nota", ["123"])


class ReportTests(unittest.TestCase):
    def test_whatsapp_flags_use_the_configured_group_and_require_opt_in(self):
        with patch(
            "sys.argv",
            ["automacao.py", "--enviar-whatsapp", "--limite", "100"],
        ):
            args = parse_args()

        self.assertTrue(args.enviar_whatsapp)
        self.assertEqual(args.grupo_whatsapp, DEFAULT_WHATSAPP_GROUP)
        self.assertEqual(
            DEFAULT_WHATSAPP_GROUP,
            "Teste Geoex",
        )
        self.assertEqual(args.limite, 100)

    def test_selects_notes_not_yet_successfully_processed(self):
        notes = ["123", "456", "789", "012", "345"]
        existing_rows = {
            "123": {"status": "Concluída"},
            "456": {"status": "Sem imagens"},
            "789": {"status": "Erro"},
            "012": {"status": "Pendente"},
            "345": {"status": "Parcial"},
        }

        self.assertEqual(
            select_notes_to_process(notes, existing_rows, 2),
            ["789", "012"],
        )

    def test_whatsapp_only_reads_existing_report_without_replacing_it(self):
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "relatorio.csv"
            write_report(report_path, ["123"], {"123": {"UTD": "CANDEIAS"}})

            with patch("automacao.write_report") as rewrite_report:
                rows = prepare_report(
                    report_path,
                    ["123"],
                    {"123": {"UTD": "OUTRO"}},
                    read_only=True,
                )

            rewrite_report.assert_not_called()
            self.assertEqual(rows["123"]["UTD"], "CANDEIAS")

    def test_report_lock_error_explains_that_excel_must_be_closed(self):
        with tempfile.TemporaryDirectory() as directory:
            report_path = Path(directory) / "relatorio.csv"
            write_report(report_path, ["123"])

            with (
                patch(
                    "automacao.os.replace",
                    side_effect=PermissionError(5, "Access is denied"),
                ),
                self.assertRaisesRegex(PermissionError, "Feche-o no Excel"),
            ):
                write_report(report_path, ["123"])

    def test_selects_next_pending_notes_after_completed_notes(self):
        notes = ["123", "456", "789"]
        existing_rows = {
            "123": {"status": "Concluída"},
            "456": {"status": "Sem imagens"},
            "789": {"status": "Pendente"},
        }

        self.assertEqual(select_notes_to_process(notes, existing_rows, 100), ["789"])

    def test_preserves_existing_results_for_notes_in_new_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "relatorio.csv"
            write_report(path, ["123", "456"])

            with path.open("r", encoding="utf-8-sig", newline="") as report_file:
                rows = list(csv.DictReader(report_file, delimiter=";"))
            rows[0]["status"] = "Concluída"
            rows[0]["motivo_nao_conclusao"] = "Acesso impedido"
            with path.open("w", encoding="utf-8-sig", newline="") as report_file:
                writer = csv.DictWriter(
                    report_file, fieldnames=rows[0].keys(), delimiter=";"
                )
                writer.writeheader()
                writer.writerows(rows)

            write_report(path, ["123", "789"])
            with path.open("r", encoding="utf-8-sig", newline="") as report_file:
                updated = list(csv.DictReader(report_file, delimiter=";"))

            self.assertEqual([row["nota"] for row in updated], ["123", "789"])
            self.assertEqual(updated[0]["status"], "Concluída")
            self.assertEqual(updated[0]["motivo_nao_conclusao"], "Acesso impedido")
            self.assertEqual(updated[1]["status"], "Pendente")

    def test_applies_salesforce_results_without_losing_other_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "relatorio.csv"
            write_report(path, ["123", "456"])
            update = {
                field: ""
                for field in REPORT_FIELDS
            }
            update.update(
                {
                    "nota": "123",
                    "UTD": "UTD-01",
                    "status": "Concluída",
                    "motivo_nao_conclusao": "Sem acesso",
                    "arquivos_baixados": "123_foto.jpg",
                }
            )
            write_report(path, ["123", "456"], {"123": update})

            with path.open("r", encoding="utf-8-sig", newline="") as report_file:
                rows = list(csv.DictReader(report_file, delimiter=";"))

            self.assertEqual(rows[0]["status"], "Concluída")
            self.assertEqual(rows[0]["UTD"], "UTD-01")
            self.assertEqual(rows[0]["motivo_nao_conclusao"], "Sem acesso")
            self.assertEqual(rows[0]["arquivos_baixados"], "123_foto.jpg")
            self.assertEqual(rows[1]["status"], "Pendente")

    def test_migrates_legacy_report_and_preserves_existing_status(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "relatorio.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as report_file:
                writer = csv.DictWriter(
                    report_file,
                    fieldnames=LEGACY_REPORT_FIELDS,
                    delimiter=";",
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "nota": "123",
                        "status": "Concluída",
                        "motivo_nao_conclusao": "Sem acesso",
                        "arquivos_baixados": "123_foto.jpg",
                        "erro": "",
                    }
                )

            write_report(
                path,
                ["123"],
                {"123": {"data_encerramento": "28/09/2026"}},
            )

            with path.open("r", encoding="utf-8-sig", newline="") as report_file:
                rows = list(csv.DictReader(report_file, delimiter=";"))

            self.assertEqual(rows[0]["status"], "Concluída")
            self.assertEqual(rows[0]["motivo_nao_conclusao"], "Sem acesso")
            self.assertEqual(rows[0]["data_encerramento"], "28/09/2026")
            self.assertEqual(rows[0]["centro_trabalho"], "")


class SalesforceLifecycleTests(unittest.TestCase):
    def test_keeps_chrome_open_when_login_does_not_finish(self):
        driver = Mock()
        with (
            patch("salesforce.criar_navegador", return_value=driver),
            patch(
                "salesforce._wait_for_login",
                side_effect=RuntimeError("login incompleto"),
            ),
        ):
            with self.assertRaisesRegex(RuntimeError, "login incompleto"):
                consultar_notas([], "https://example.my.salesforce.com/", Path("."), None)

        driver.service.stop.assert_called_once()
        driver.quit.assert_not_called()

    def test_closes_chrome_after_successful_authenticated_run(self):
        driver = Mock()
        on_result = Mock()
        results_to_process = [
            ResultadoNota(nota="123", status="Concluída"),
            ResultadoNota(nota="456", status="Sem imagens"),
        ]
        with (
            patch("salesforce.criar_navegador", return_value=driver),
            patch("salesforce._wait_for_login"),
            patch(
                "salesforce._processar_nota",
                side_effect=results_to_process,
            ),
        ):
            results = consultar_notas(
                ["123", "456"],
                "https://example.my.salesforce.com/",
                Path("."),
                None,
                on_result=on_result,
            )

        self.assertEqual(
            [result.status for result in results],
            ["Concluída", "Sem imagens"],
        )
        self.assertEqual(
            [call.args[0] for call in on_result.call_args_list],
            results,
        )
        driver.quit.assert_called_once()
        driver.service.stop.assert_not_called()

    def test_chrome_uses_a_detached_profile_for_manual_login(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            driver = Mock()
            with (
                patch.dict("os.environ", {"LOCALAPPDATA": str(root)}),
                patch("salesforce.webdriver.Chrome", return_value=driver) as chrome,
            ):
                criar_navegador(root / "downloads", root / "chrome-profile")

            options = chrome.call_args.kwargs["options"]
            self.assertTrue(options.experimental_options["detach"])


class SalesforceNavigationTests(unittest.TestCase):
    def test_search_result_is_activated_with_javascript_click(self):
        search_button = Mock()
        search_box = Mock()
        search_result = Mock()
        first_wait = Mock()
        first_wait.until.return_value = search_button
        result_wait = Mock()
        result_wait.until.return_value = search_result
        navigation_wait = Mock()
        navigation_wait.until.side_effect = lambda condition: condition(driver)
        driver = Mock()
        driver.current_url = "https://example.lightning.force.com/lightning/r/0WOR123/view"
        driver.find_elements.return_value = [
            SimpleNamespace(text="Ordem de Serviço COE_CS1030177305")
        ]

        with (
            patch(
                "salesforce.WebDriverWait",
                side_effect=[result_wait, navigation_wait],
            ),
            patch(
                "salesforce._wait_for_stable_search_box",
                side_effect=[False, False, search_box],
            ),
        ):
            _buscar_ordem(driver, first_wait, "1030177305")

        search_button.click.assert_called_once()
        search_box.click.assert_not_called()
        search_box.send_keys.assert_called_once_with("1030177305")
        self.assertEqual(driver.execute_script.call_count, 2)
        self.assertIs(driver.execute_script.call_args_list[0].args[1], search_button)
        self.assertIs(driver.execute_script.call_args_list[1].args[1], search_result)
        search_result.click.assert_not_called()

    def test_opens_related_tab_and_waits_for_it_to_be_selected(self):
        selected = {"value": False}
        tab = Mock()
        tab.text = "Relacionado"
        tab.get_attribute.side_effect = lambda name: (
            "true" if selected["value"] else "false"
        ) if name == "aria-selected" else None
        download_button = Mock()
        download_button.is_displayed.return_value = True
        download_button.is_enabled.return_value = True
        driver = Mock()
        driver.find_elements.side_effect = [[tab], [tab]]

        def click_related_tab(_script, clicked_tab):
            self.assertIs(clicked_tab, tab)
            selected["value"] = True

        driver.execute_script.side_effect = click_related_tab
        wait = Mock()
        wait.until.side_effect = lambda condition: condition(driver)

        _abrir_aba_relacionados(driver, wait, "1030177305")

        self.assertTrue(selected["value"])
        self.assertEqual(driver.find_elements.call_count, 2)

    def test_finds_service_appointments_in_the_work_order_related_card(self):
        appointment_url = (
            "https://example.lightning.force.com/"
            "lightning/r/08pRo000000000001/view"
        )
        appointment_link = Mock()
        appointment_link.get_attribute.return_value = appointment_url
        card = Mock()
        card.find_elements.side_effect = lambda _by, selector: (
            [appointment_link]
            if selector == ".//a[contains(@href, '/lightning/r/08p')]"
            else []
        )
        driver = Mock()
        driver.find_elements.return_value = [card]
        wait = Mock()
        wait.until.side_effect = lambda condition: condition(driver)

        result = _compromissos_relacionados(wait, "1030177305")

        self.assertEqual(result, [appointment_url])

    def test_returns_empty_when_related_work_order_has_no_appointments(self):
        heading = SimpleNamespace(text="Compromissos de serviços (0)")
        card = Mock()
        card.find_elements.side_effect = lambda _by, selector: (
            [heading]
            if selector == ".//h2[contains(., 'Compromissos de serviços')]"
            else []
        )
        driver = Mock()
        driver.find_elements.return_value = [card]
        wait = Mock()
        wait.until.side_effect = lambda condition: condition(driver)

        self.assertEqual(
            _compromissos_relacionados(wait, "1030177305"),
            [],
        )

    def test_opens_service_appointment_from_the_related_card(self):
        appointment_id = "08pRo000000000001"
        appointment_url = (
            f"https://example.lightning.force.com/lightning/r/{appointment_id}/view"
        )
        driver = Mock()
        driver.current_url = (
            "https://example.lightning.force.com/lightning/r/WorkOrder/123/view"
        )

        def navigate_to_appointment(url):
            self.assertEqual(url, appointment_url)
            driver.current_url = (
                "https://example.lightning.force.com/lightning/r/"
                f"ServiceAppointment/{appointment_id}/view"
            )

        driver.get.side_effect = navigate_to_appointment
        wait = Mock()
        wait.until.side_effect = lambda condition: condition(driver)

        _abrir_compromisso_servico(
            driver,
            wait,
            "1030177305",
            appointment_url,
        )

        self.assertIn(appointment_id, driver.current_url)
        driver.get.assert_called_once_with(appointment_url)

    def test_search_box_searches_open_shadow_roots(self):
        visible_input = Mock()
        driver = Mock()
        driver.execute_script.return_value = visible_input

        self.assertIs(_visible_search_box(driver), visible_input)
        script = driver.execute_script.call_args.args[0]
        self.assertIn("element.shadowRoot", script)
        self.assertIn("type === 'search'", script)
        self.assertIn("rect.width > 0", script)
        self.assertIn("document.elementFromPoint", script)
        self.assertIn("topElement.shadowRoot.elementFromPoint", script)
        self.assertIn("return topElement === element", script)
        self.assertIn("isAccessibleFromPage(element)", script)

    def test_search_box_waits_for_stable_geometry(self):
        search_box = Mock()
        driver = Mock()
        driver.execute_script.side_effect = [
            [0, 0, 0, 0],
            [321, 52, 664, 32],
            [321, 52, 664, 32],
            [321, 52, 664, 32],
        ]

        def until_stable(condition):
            for _ in range(4):
                result = condition(driver)
                if result:
                    return result
            raise TimeoutException()

        with (
            patch("salesforce.WebDriverWait") as wait_class,
            patch(
                "salesforce._visible_search_box",
                return_value=search_box,
            ) as visible_search_box,
        ):
            wait_class.return_value.until.side_effect = until_stable

            result = _wait_for_stable_search_box(driver, timeout=2)

        self.assertIs(result, search_box)
        self.assertEqual(visible_search_box.call_count, 4)

    def test_search_box_does_not_select_input_outside_dialog(self):
        driver = Mock()
        driver.execute_script.return_value = None

        self.assertFalse(_visible_search_box(driver))

    def test_search_box_excludes_salesforce_object_type_picker(self):
        driver = Mock()
        driver.execute_script.return_value = None

        self.assertFalse(_visible_search_box(driver))
        self.assertIn(
            "by object type",
            driver.execute_script.call_args.args[0],
        )

    def test_ignores_hidden_search_results(self):
        hidden = SimpleNamespace(
            is_displayed=lambda: False,
            is_enabled=lambda: True,
            text="04596443 Ordem de trabalho • COE_CS1030177305",
        )
        visible = SimpleNamespace(
            is_displayed=lambda: True,
            is_enabled=lambda: True,
            text="04596443 Ordem de trabalho • COE_CS1030177305",
        )
        driver = Mock()
        driver.find_elements.return_value = [hidden, visible]

        self.assertIs(
            _matching_work_order_result(driver, "1030177305"),
            visible,
        )

    def test_extracts_all_images_from_zip_and_ignores_other_files(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            archive_path = destination / "download.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("equipe/foto1.JPG", b"image one")
                archive.writestr("equipe/foto2.png", b"image two")
                archive.writestr("equipe/relatorio.pdf", b"document")

            names = _extract_downloaded_images(
                [archive_path],
                "1030177305",
                destination,
            )

            self.assertEqual(
                names,
                ["1030177305_foto1.JPG", "1030177305_foto2.png"],
            )
            self.assertEqual(
                (destination / names[0]).read_bytes(),
                b"image one",
            )
            self.assertEqual(
                (destination / names[1]).read_bytes(),
                b"image two",
            )
            self.assertFalse(archive_path.exists())


class WhatsAppDeliveryTests(unittest.TestCase):
    def test_document_send_waits_for_new_message_to_be_marked_sent(self):
        driver = Mock()
        archive_path = Path("Respostas.zip")
        old_message = {"id": "old", "sent": True}
        pending_message = {"id": "new", "sent": False}
        sent_message = {"id": "new", "sent": True}
        wait_timeouts = []

        class ImmediateWait:
            def __init__(self, _driver, timeout):
                wait_timeouts.append(timeout)

            def until(self, condition):
                for _ in range(3):
                    if condition(driver):
                        return True
                raise AssertionError("The document was not confirmed as sent.")

        with (
            patch(
                "whatsapp._media_dialog",
                side_effect=[Mock(), False, False],
            ),
            patch(
                "whatsapp._document_message_states",
                side_effect=[
                    [old_message, pending_message],
                    [old_message, sent_message],
                ],
            ),
            patch("whatsapp.WebDriverWait", ImmediateWait),
        ):
            _wait_for_document_sent(driver, archive_path, {"old"})

        self.assertEqual(wait_timeouts, [300])

    def test_document_send_timeout_is_not_treated_as_success(self):
        driver = Mock()
        archive_path = Path("Respostas.zip")

        class ImmediateWait:
            def __init__(self, _driver, _timeout):
                pass

            def until(self, condition):
                if condition(driver):
                    return True
                raise TimeoutException()

        with (
            patch("whatsapp._media_dialog", return_value=False),
            patch(
                "whatsapp._document_message_states",
                return_value=[{"id": "new", "sent": False}],
            ),
            patch("whatsapp.WebDriverWait", ImmediateWait),
        ):
            with self.assertRaisesRegex(
                RuntimeError, "não confirmou o envio"
            ):
                _wait_for_document_sent(driver, archive_path, set())

    def test_legacy_delivery_state_retries_unconfirmed_archive(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output_dir = root / "Respostas"
            output_dir.mkdir()
            state_path = root / ".Respostas_whatsapp_state.json"
            state_path.write_text(
                json.dumps({"notes": {}, "archive_sha256": "old-hash"}),
                encoding="utf-8",
            )
            driver = Mock()

            with (
                patch("whatsapp._create_browser", return_value=driver),
                patch("whatsapp._wait_for_search_input"),
                patch("whatsapp._open_group"),
                patch("whatsapp._send_document") as send_document,
            ):
                enviar_respostas(
                    output_dir,
                    {},
                    DEFAULT_WHATSAPP_GROUP,
                    input_fn=lambda _prompt: "s",
                )

            send_document.assert_called_once_with(
                driver, root / "Respostas.zip"
            )
            state = json.loads(state_path.read_text(encoding="utf-8"))
            self.assertEqual(state["version"], 2)
            self.assertTrue(state["archive_sha256"])

    def test_retries_fresh_attachment_button_lookup_after_stale_dom(self):
        driver = Mock()
        driver.execute_script.side_effect = [
            StaleElementReferenceException(),
            False,
            True,
        ]

        class ImmediateWait:
            def __init__(self, _driver, _timeout):
                pass

            def until(self, condition):
                for _ in range(3):
                    if condition(driver):
                        return True
                raise AssertionError("Expected attachment button to appear.")

        with patch("whatsapp.WebDriverWait", ImmediateWait):
            _click_attachment_button(driver)

        self.assertEqual(driver.execute_script.call_count, 3)
        self.assertTrue(
            all(
                "button.click()" in call.args[0] and len(call.args) == 1
                for call in driver.execute_script.call_args_list
            )
        )

    def test_selects_media_menu_item_with_fresh_dom_lookup(self):
        driver = Mock()
        driver.execute_script.side_effect = [
            StaleElementReferenceException(),
            False,
            True,
        ]

        class ImmediateWait:
            def __init__(self, _driver, _timeout):
                pass

            def until(self, condition):
                for _ in range(3):
                    if condition(driver):
                        return True
                raise AssertionError("Expected menu item to become available.")

        with (
            patch("whatsapp.WebDriverWait", ImmediateWait),
        ):
            _click_media_menu_item(driver)

        self.assertEqual(driver.execute_script.call_count, 3)
        self.assertTrue(
            all(
                "controls[0].click()" in call.args[0]
                and "text.includes(label)" in call.args[0]
                and len(call.args) == 1
                for call in driver.execute_script.call_args_list
            )
        )

    def test_media_menu_timeout_is_actionable(self):
        driver = Mock()
        driver.execute_script.return_value = False

        class ImmediateWait:
            def __init__(self, _driver, _timeout):
                pass

            def until(self, condition):
                for _ in range(2):
                    condition(driver)
                raise TimeoutException()

        with (
            patch("whatsapp.WebDriverWait", ImmediateWait),
        ):
            with self.assertRaisesRegex(
                TimeoutException, "não apareceu ou não pôde ser selecionada"
            ) as raised:
                _click_media_menu_item(driver)

        self.assertEqual(driver.execute_script.call_count, 3)
        self.assertIn("Elementos visíveis", str(raised.exception))

    def test_photo_upload_prefers_image_and_video_input_over_image_only(self):
        image_input = Mock()
        image_input.get_attribute.return_value = "image/*"
        photo_video_input = Mock()
        photo_video_input.get_attribute.return_value = (
            "image/*,video/mp4,video/3gpp,video/quicktime"
        )
        driver = Mock()
        driver.find_elements.return_value = [image_input, photo_video_input]

        self.assertIs(
            _find_file_input(driver, images_only=True),
            photo_video_input,
        )

    def test_photo_upload_falls_back_when_only_image_input_exists(self):
        image_input = Mock()
        image_input.get_attribute.return_value = "image/*"
        driver = Mock()
        driver.find_elements.return_value = [image_input]

        self.assertIs(_find_file_input(driver, images_only=True), image_input)

    def test_archive_contains_the_entire_responses_folder_and_is_deterministic(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output_dir = root / "Respostas"
            note_dir = output_dir / "123"
            note_dir.mkdir(parents=True)
            (output_dir / "relatorio_processamento.csv").write_text(
                "nota;status\n123;Concluída\n",
                encoding="utf-8",
            )
            (note_dir / "descricao.txt").write_text("Nota: 123\n", encoding="utf-8")
            (note_dir / "123_foto.jpg").write_bytes(b"image")

            archive_path = create_responses_archive(output_dir)
            first_archive = archive_path.read_bytes()
            create_responses_archive(output_dir)

            self.assertEqual(archive_path, root / "Respostas.zip")
            self.assertEqual(archive_path.read_bytes(), first_archive)
            with zipfile.ZipFile(archive_path) as archive:
                self.assertEqual(
                    set(archive.namelist()),
                    {
                        "relatorio_processamento.csv",
                        "123/descricao.txt",
                        "123/123_foto.jpg",
                    },
                )

    def test_note_message_includes_salesforce_impediment_reason(self):
        with tempfile.TemporaryDirectory() as directory:
            description = Path(directory) / "descricao.txt"
            description.write_text("Nota: 123\nCódigo de ação: NE08\n", encoding="utf-8")

            message = _note_message(
                description,
                {"motivo_nao_conclusao": "Conta paga"},
            )

        self.assertEqual(
            message,
            "Nota: 123\nCódigo de ação: NE08\nMotivo do impedimento: Conta paga",
        )

    def test_group_name_is_escaped_as_an_xpath_literal(self):
        self.assertEqual(
            _xpath_literal("""Grupo "A" d'água"""),
            """concat('Grupo "A" d', "'", 'água')""",
        )

    def test_group_search_falls_back_to_focusing_when_click_is_intercepted(self):
        group_name = "Pendências Geoex - Metropolitana"
        search_input = Mock()
        search_input.click.side_effect = ElementClickInterceptedException(
            "A search input is covered."
        )
        result = Mock()
        result.is_displayed.return_value = True
        result.click.side_effect = ElementClickInterceptedException(
            "The row container should be clicked instead."
        )
        row = Mock()
        row.is_displayed.return_value = True
        row.click.side_effect = ElementClickInterceptedException(
            "Use JavaScript on the result row."
        )
        header = Mock()
        header.get_attribute.return_value = group_name
        driver = Mock()

        def find_elements(by, selector):
            if selector.startswith("//*["):
                return [result]
            if selector.startswith("header "):
                return [header]
            return []

        driver.find_elements.side_effect = find_elements
        driver.execute_script.side_effect = lambda script, *args: (
            row if "closest(" in script else args[0] if args else None
        )
        with patch("whatsapp._wait_for_search_input", return_value=search_input):
            _open_group(driver, group_name)

        driver.execute_script.assert_any_call("arguments[0].focus()", search_input)
        search_input.send_keys.assert_any_call(group_name)
        row.click.assert_called_once()
        self.assertTrue(
            any(
                "arguments[0].click();" in call.args[0]
                and call.args[1] is row
                for call in driver.execute_script.call_args_list
            )
        )
        self.assertIn(
            "contains(@aria-label,",
            driver.find_elements.call_args_list[0].args[1],
        )

    def test_group_search_accepts_accessible_name_and_header_heading(self):
        group_name = "Pendências Geoex - Metropolitana"
        search_input = Mock()
        result = Mock()
        result.is_displayed.return_value = True
        result.get_attribute.side_effect = lambda name: (
            group_name if name == "aria-label" else None
        )
        header_heading = Mock()
        header_heading.is_displayed.return_value = True
        header_heading.get_attribute.return_value = None
        header_heading.text = group_name
        driver = Mock()

        def find_elements(by, selector):
            if selector.startswith("//*["):
                return [result]
            if selector.startswith("header "):
                return [header_heading]
            return []

        driver.find_elements.side_effect = find_elements
        driver.execute_script.side_effect = (
            lambda _script, *args: args[0] if args else None
        )
        with patch("whatsapp._wait_for_search_input", return_value=search_input):
            _open_group(driver, group_name)

        result.click.assert_called_once()
        self.assertTrue(driver.execute_script.called)

    def test_group_search_matches_result_with_message_preview_text(self):
        group_name = "Teste Geoex"
        search_input = Mock()
        result = Mock()
        result.is_displayed.return_value = True
        result.get_attribute.return_value = None
        result.text = "Teste Geoex\nVocê: teste de validação\n19:30"
        header = Mock()
        header.is_displayed.return_value = True
        header.get_attribute.return_value = None
        header.text = group_name
        driver = Mock()

        def find_elements(by, selector):
            if selector.startswith("//*"):
                self.assertIn("contains(normalize-space(.), 'Teste Geoex')", selector)
                return [result]
            if selector.startswith("header "):
                return [header]
            return []

        driver.find_elements.side_effect = find_elements
        driver.execute_script.side_effect = (
            lambda _script, *args: args[0] if args else None
        )
        with patch("whatsapp._wait_for_search_input", return_value=search_input):
            _open_group(driver, group_name)

        result.click.assert_called_once()

    def test_sends_all_images_together_with_one_caption(self):
        images = [Path("one.jpg"), Path("two.jpg"), Path("three.jpg")]
        driver = Mock()
        with patch("whatsapp._send_images_with_caption") as send_images:
            _send_media(driver, images, "Descrição da nota")

        send_images.assert_called_once_with(
            driver, images, "Descrição da nota"
        )

    def test_media_preview_is_detected_by_caption_when_dialog_role_is_absent(self):
        caption_box = Mock()
        caption_box.is_displayed.return_value = True
        caption_box.get_attribute.side_effect = lambda name: (
            "Adicione uma legenda"
            if name == "data-placeholder"
            else None
        )
        preview_container = Mock()
        preview_container.is_displayed.return_value = True
        driver = Mock()
        driver.find_elements.side_effect = lambda by, selector: (
            []
            if selector == "//*[@role='dialog']"
            or "media-editor" in selector
            else [caption_box]
        )
        driver.execute_script.return_value = preview_container

        self.assertIs(_media_dialog(driver), preview_container)
        self.assertEqual(
            driver.execute_script.call_args.args[1],
            caption_box,
        )

    def test_media_preview_falls_back_to_send_button_container(self):
        preview_container = Mock()
        preview_container.is_displayed.return_value = True
        caption_box = Mock()
        caption_box.is_displayed.return_value = True
        caption_box.get_attribute.return_value = None
        driver = Mock()
        driver.find_elements.side_effect = lambda by, selector: (
            []
            if selector == "//*[@role='dialog']"
            or "media-editor" in selector
            else [caption_box]
        )
        driver.execute_script.return_value = preview_container

        self.assertIs(_media_dialog(driver), preview_container)
        self.assertEqual(len(driver.execute_script.call_args.args), 1)
        self.assertIn(
            "data-testid*=\"send\"",
            driver.execute_script.call_args.args[0],
        )

    def test_media_preview_is_detected_by_media_editor_canvas(self):
        canvas = Mock()
        canvas.is_displayed.return_value = True
        canvas.get_attribute.return_value = "media-editor-canvas"
        editor_container = Mock()
        editor_container.is_displayed.return_value = True
        driver = Mock()
        driver.find_elements.side_effect = lambda by, selector: (
            [canvas]
            if "media-editor" in selector
            else []
        )
        driver.execute_script.return_value = editor_container

        self.assertIs(_media_dialog(driver), editor_container)
        self.assertEqual(
            driver.execute_script.call_args.args[1],
            canvas,
        )
        self.assertIn(
            "node.querySelectorAll",
            driver.execute_script.call_args.args[0],
        )

    def test_media_send_button_is_recognized_by_nested_send_icon(self):
        send_button = Mock()
        send_button.text = ""
        send_button.get_attribute.return_value = None
        send_button.is_displayed.return_value = True
        send_icon = Mock()
        send_icon.get_attribute.return_value = "send"
        send_button.find_elements.return_value = [send_icon]
        preview_container = Mock()
        preview_container.find_elements.return_value = [send_button]

        self.assertIs(_send_button(preview_container), send_button)

    def test_multiple_photos_use_photo_video_menu_and_send_one_caption(self):
        with tempfile.TemporaryDirectory() as directory:
            images = [
                Path(directory) / "foto-1.jpg",
                Path(directory) / "foto-2.jpg",
            ]
            for image in images:
                image.write_bytes(b"original image")
            driver = Mock()
            file_input = Mock()
            file_input.get_property.return_value = True
            dialog = Mock()
            dialog.is_displayed.return_value = True
            caption_box = Mock()
            caption_box.is_displayed.return_value = True
            dialog.find_elements.return_value = [caption_box]
            send_button = Mock()
            media_dialog_calls = iter([dialog, dialog, dialog, False])

            with (
                patch("whatsapp._click_attachment_button"),
                patch("whatsapp._wait_for_conversation_compose_box"),
                patch("whatsapp._click_media_menu_item"),
                patch("whatsapp._find_file_input", return_value=file_input),
                patch("whatsapp._conversation_compose_box", return_value=Mock()),
                patch(
                    "whatsapp._media_dialog",
                    side_effect=lambda _driver: next(media_dialog_calls),
                ),
                patch("whatsapp._send_button", return_value=send_button),
                patch(
                    "whatsapp.WebDriverWait",
                    return_value=Mock(
                        until=lambda condition: condition(driver)
                    ),
                ),
            ):
                _send_images_with_caption(
                    driver, images, "Primeira linha\nSegunda linha"
                )

        self.assertEqual(
            [
                call
                for call in driver.execute_script.call_args_list
                if call.args[0] == "arguments[0].click()"
            ],
            [unittest.mock.call("arguments[0].click()", send_button)],
        )
        file_input.send_keys.assert_called_once_with(
            "\n".join(str(image.resolve()) for image in images)
        )
        self.assertEqual(
            caption_box.send_keys.call_args_list,
            [
                unittest.mock.call("Primeira linha"),
                unittest.mock.call(Keys.SHIFT, Keys.ENTER),
                unittest.mock.call("Segunda linha"),
            ],
        )
        self.assertEqual(
            len(
                [
                    call
                    for call in driver.execute_script.call_args_list
                    if call.args[0] == "arguments[0].click()"
                ]
            ),
            1,
        )

    def test_multiple_photos_fail_before_upload_if_input_lacks_multiple(self):
        with tempfile.TemporaryDirectory() as directory:
            images = [Path(directory) / "one.jpg", Path(directory) / "two.jpg"]
            driver = Mock()
            file_input = Mock()
            file_input.get_property.return_value = False
            with (
                patch("whatsapp._click_attachment_button"),
                patch("whatsapp._wait_for_conversation_compose_box"),
                patch("whatsapp._click_media_menu_item"),
                patch("whatsapp._find_file_input", return_value=file_input),
            ):
                with self.assertRaisesRegex(
                    RuntimeError, "não aceita múltiplos arquivos"
                ):
                    _send_images_with_caption(
                        driver, images, "Descrição da nota"
                    )

        file_input.send_keys.assert_not_called()
        self.assertFalse(
            any(
                call.args[0] == "arguments[0].click()"
                for call in driver.execute_script.call_args_list
            )
        )

    def test_missing_caption_fails_without_sending_media_or_separate_text(self):
        with tempfile.TemporaryDirectory() as directory:
            image = Path(directory) / "foto.jpg"
            image.write_bytes(b"original image")
            driver = Mock()
            file_input = Mock()
            file_input.get_property.return_value = True
            media_editor = Mock()
            media_editor.find_elements.return_value = []

            class ImmediateWait:
                def __init__(self, _driver, _timeout):
                    pass

                def until(self, condition):
                    result = condition(driver)
                    if not result:
                        raise TimeoutException()
                    return result

            with (
                patch("whatsapp._click_attachment_button"),
                patch("whatsapp._wait_for_conversation_compose_box"),
                patch("whatsapp._click_media_menu_item"),
                patch("whatsapp._find_file_input", return_value=file_input),
                patch("whatsapp._media_dialog", return_value=media_editor),
                patch("whatsapp._media_preview_diagnostics", return_value="{}"),
                patch("whatsapp.WebDriverWait", ImmediateWait),
            ):
                with self.assertRaisesRegex(
                    RuntimeError, "campo da descrição"
                ):
                    _send_images_with_caption(
                        driver, [image], "Descrição da nota"
                    )

        self.assertFalse(
            any(
                call.args[0] == "arguments[0].click()"
                for call in driver.execute_script.call_args_list
            )
        )

    def test_sends_images_then_full_zip_and_skips_identical_repeat(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output_dir = root / "Respostas"
            note_dir = output_dir / "123"
            note_dir.mkdir(parents=True)
            (output_dir / "relatorio_processamento.csv").write_text(
                "nota;status\n123;Concluída\n456;Sem imagens\n",
                encoding="utf-8",
            )
            (note_dir / "descricao.txt").write_text("Nota: 123\n", encoding="utf-8")
            (note_dir / "123_foto.jpg").write_bytes(b"image data")
            empty_note_dir = output_dir / "456"
            empty_note_dir.mkdir()
            (empty_note_dir / "descricao.txt").write_text(
                "Nota: 456\n",
                encoding="utf-8",
            )
            driver = Mock()
            report_rows = {
                "123": {
                    "status": "Concluída",
                    "motivo_nao_conclusao": "Conta paga",
                },
                "456": {
                    "status": "Sem imagens",
                    "motivo_nao_conclusao": "",
                },
                "789": {
                    "status": "Erro",
                    "motivo_nao_conclusao": "",
                },
            }

            with (
                patch("whatsapp._create_browser", return_value=driver),
                patch("whatsapp._wait_for_search_input"),
                patch("whatsapp._open_group") as open_group,
                patch("whatsapp._send_media") as send_media,
                patch("whatsapp._send_document") as send_document,
            ):
                confirmation_prompts = []
                enviar_respostas(
                    output_dir,
                    report_rows,
                    DEFAULT_WHATSAPP_GROUP,
                    input_fn=lambda prompt: (
                        confirmation_prompts.append(prompt) or "s"
                    ),
                )

            self.assertEqual(
                confirmation_prompts,
                [
                    "Confirme que o grupo aberto é "
                    f"'{DEFAULT_WHATSAPP_GROUP}' e autorize os envios [s/N]: "
                ],
            )
            open_group.assert_called_once_with(driver, DEFAULT_WHATSAPP_GROUP)
            send_media.assert_called_once()
            media_args = send_media.call_args.args
            self.assertEqual(media_args[0], driver)
            self.assertEqual(media_args[1], [note_dir / "123_foto.jpg"])
            self.assertIn("Motivo do impedimento: Conta paga", media_args[2])
            send_document.assert_called_once_with(driver, root / "Respostas.zip")
            driver.quit.assert_called_once()
            with zipfile.ZipFile(root / "Respostas.zip") as archive:
                self.assertIn("123/123_foto.jpg", archive.namelist())
                self.assertIn("456/descricao.txt", archive.namelist())
                self.assertIn(
                    "456",
                    archive.read("relatorio_processamento.csv").decode(),
                )

            with patch("whatsapp._create_browser") as create_browser:
                enviar_respostas(
                    output_dir,
                    report_rows,
                    DEFAULT_WHATSAPP_GROUP,
                    input_fn=lambda _prompt: self.fail(
                        "Already delivered content must not request confirmation."
                    ),
                )

            create_browser.assert_not_called()

    def test_resumes_after_partial_failure_without_duplicating_sent_notes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output_dir = root / "Respostas"
            report_rows = {}
            for note in ("123", "456"):
                note_dir = output_dir / note
                note_dir.mkdir(parents=True)
                (note_dir / "descricao.txt").write_text(
                    f"Nota: {note}\n",
                    encoding="utf-8",
                )
                (note_dir / f"{note}_foto.jpg").write_bytes(
                    f"image {note}".encode()
                )
                report_rows[note] = {
                    "status": "Concluída",
                    "motivo_nao_conclusao": "",
                }

            first_driver = Mock()
            resume_driver = Mock()
            with (
                patch(
                    "whatsapp._create_browser",
                    return_value=first_driver,
                ),
                patch("whatsapp._wait_for_search_input"),
                patch("whatsapp._open_group"),
                patch(
                    "whatsapp._send_media",
                    side_effect=[None, RuntimeError("falha simulada")],
                ) as first_run_send,
                patch("whatsapp._send_document"),
            ):
                with self.assertRaisesRegex(RuntimeError, "falha simulada"):
                    enviar_respostas(
                        output_dir,
                        report_rows,
                        DEFAULT_WHATSAPP_GROUP,
                        input_fn=lambda _prompt: "s",
                    )

            self.assertEqual(first_run_send.call_count, 2)
            first_driver.quit.assert_called_once()

            with (
                patch(
                    "whatsapp._create_browser",
                    return_value=resume_driver,
                ),
                patch("whatsapp._wait_for_search_input"),
                patch("whatsapp._open_group"),
                patch("whatsapp._send_media") as resume_send,
                patch("whatsapp._send_document") as send_document,
            ):
                enviar_respostas(
                    output_dir,
                    report_rows,
                    DEFAULT_WHATSAPP_GROUP,
                    input_fn=lambda _prompt: "s",
                )

            self.assertEqual(resume_send.call_count, 1)
            self.assertEqual(
                resume_send.call_args.args[1],
                [output_dir / "456" / "456_foto.jpg"],
            )
            send_document.assert_called_once_with(
                resume_driver, root / "Respostas.zip"
            )
            resume_driver.quit.assert_called_once()

            with patch("whatsapp._create_browser") as create_browser:
                enviar_respostas(
                    output_dir,
                    report_rows,
                    DEFAULT_WHATSAPP_GROUP,
                    input_fn=lambda _prompt: self.fail(
                        "Identical notes and ZIP must not be sent again."
                    ),
                )

            create_browser.assert_not_called()

class SalesforceDownloadTests(unittest.TestCase):
    def test_empty_zip_returns_no_images_without_failing(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            archive_path = destination / "download.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("readme.txt", "no photos")

            self.assertEqual(
                _extract_downloaded_images(
                    [archive_path],
                    "1030177305",
                    destination,
                ),
                [],
            )
            self.assertFalse(archive_path.exists())

    def test_corrupt_zip_is_reported_instead_of_marked_as_no_images(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory)
            archive_path = destination / "download.zip"
            archive_path.write_bytes(b"not a zip archive")

            with self.assertRaisesRegex(RuntimeError, "Não foi possível extrair"):
                _extract_downloaded_images(
                    [archive_path],
                    "1030177305",
                    destination,
                )

            self.assertTrue(archive_path.exists())

    def test_empty_attachments_are_reported_as_no_images(self):
        with tempfile.TemporaryDirectory() as directory:
            driver = Mock()
            driver.current_url = (
                "https://example.lightning.force.com/lightning/r/WorkOrder/123/view"
            )
            with (
                patch("salesforce.WebDriverWait", return_value=Mock()),
                patch("salesforce._buscar_ordem"),
                patch("salesforce._abrir_aba_relacionados"),
                patch(
                    "salesforce._compromissos_relacionados",
                    return_value=[
                        "https://example.lightning.force.com/lightning/r/"
                        "08pRo000000000001/view"
                    ],
                ),
                patch("salesforce._abrir_compromisso_servico"),
                patch("salesforce._ler_motivo", return_value=""),
                patch("salesforce._baixar_arquivos_relacionados", return_value=[]),
            ):
                result = _processar_nota(
                    driver,
                    Path(directory),
                    "1030177305",
                )

            self.assertEqual(result.status, "Sem imagens")
            self.assertEqual(result.arquivos, [])
            self.assertEqual(result.erro, "")

    def test_opens_related_service_appointment_before_downloading_files(self):
        with tempfile.TemporaryDirectory() as directory:
            driver = Mock()
            driver.current_url = (
                "https://example.lightning.force.com/lightning/r/WorkOrder/123/view"
            )
            events = []
            with (
                patch("salesforce.WebDriverWait", return_value=Mock()),
                patch("salesforce._buscar_ordem"),
                patch(
                    "salesforce._abrir_aba_relacionados",
                    side_effect=lambda *_args: events.append("related"),
                ) as open_related,
                patch(
                    "salesforce._compromissos_relacionados",
                    return_value=[
                        "https://example.lightning.force.com/lightning/r/"
                        "08pRo000000000001/view"
                    ],
                ),
                patch(
                    "salesforce._abrir_compromisso_servico",
                    side_effect=lambda *_args: events.append("appointment"),
                ) as open_appointment,
                patch("salesforce._ler_motivo", return_value=""),
                patch(
                    "salesforce._baixar_arquivos_relacionados",
                    side_effect=lambda *_args: events.append("download")
                    or ["1030177305_foto.jpg"],
                ) as download_files,
            ):
                result = _processar_nota(
                    driver,
                    Path(directory),
                    "1030177305",
                )

            self.assertEqual(result.status, "Concluída")
            self.assertEqual(result.arquivos, ["1030177305_foto.jpg"])
            self.assertEqual(
                events,
                ["related", "appointment", "download"],
            )
            open_related.assert_called_once()
            open_appointment.assert_called_once()
            download_files.assert_called_once()
            driver.get.assert_not_called()

    def test_empty_file_modal_returns_no_images_without_clicking_download(self):
        toolbar_button = Mock()
        close_button = Mock()
        modal = Mock()
        modal.is_displayed.return_value = True
        modal.text = "Nenhum arquivo"

        def modal_elements(_by, selector):
            if selector == ".//button[contains(., 'Cancelar e fechar')]":
                return [close_button]
            return []

        modal.find_elements.side_effect = modal_elements

        initial_wait = Mock()
        initial_wait.until.return_value = toolbar_button
        modal_wait = Mock()
        modal_wait.until.side_effect = lambda condition: condition(driver)
        file_list_wait = Mock()
        file_list_wait.until.side_effect = lambda condition: condition(driver)
        driver = Mock()
        driver.find_elements.return_value = [modal]
        driver.execute_script.return_value = []

        with (
            tempfile.TemporaryDirectory() as directory,
            patch(
                "salesforce.WebDriverWait",
                side_effect=[modal_wait, file_list_wait],
            ),
        ):
            result = _baixar_arquivos_relacionados(
                driver,
                initial_wait,
                "1030177305",
                Path(directory),
            )
        self.assertEqual(result, [])
        toolbar_button.click.assert_called_once()
        driver.execute_script.assert_any_call(
            "arguments[0].click()",
            close_button,
        )
        self.assertEqual(driver.execute_script.call_count, 3)

    def test_modal_image_rows_are_found_inside_shadow_dom(self):
        checkbox = Mock()
        modal = Mock()
        driver = Mock()
        driver.execute_script.return_value = [
            {"checkbox": checkbox, "filename": "10001129.jpg"},
            {"checkbox": Mock(), "filename": "10001130.jpg"},
        ]

        rows = _modal_image_rows(driver, modal)

        self.assertEqual(
            rows,
            [(checkbox, "10001129.jpg"), (driver.execute_script.return_value[1]["checkbox"], "10001130.jpg")],
        )
        script = driver.execute_script.call_args.args[0]
        self.assertIn("element.shadowRoot", script)
        self.assertIn("input[type=\"checkbox\"]", script)

    def test_modal_download_button_is_found_inside_shadow_dom(self):
        modal = Mock()
        button = Mock()
        driver = Mock()
        driver.execute_script.return_value = button

        result = _modal_download_button(driver, modal)

        self.assertIs(result, button)
        script = driver.execute_script.call_args.args[0]
        self.assertIn("element.shadowRoot", script)
        self.assertIn("label === 'baixar'", script)

    def test_waits_for_modal_file_rows_before_selecting_and_downloading(self):
        heading = Mock()
        heading.is_displayed.return_value = True
        toolbar_button = Mock()
        checkbox = Mock()
        checkbox.is_selected.return_value = False
        photo_rows = [{"checkbox": checkbox, "filename": "equipe/foto.jpg"}]
        modal = Mock()
        modal.is_displayed.return_value = True
        modal.text = "Arquivo Tipo Baixar"
        modal.find_elements.side_effect = lambda _by, selector: (
            [download_button]
            if selector == ".//button[normalize-space(.)='Baixar']"
            else []
        )
        download_button = Mock()

        def create_zip():
            with zipfile.ZipFile(Path(directory) / "download.zip", "w") as archive:
                archive.writestr("equipe/foto.jpg", b"image")
                archive.writestr("equipe/relatorio.pdf", b"document")

        initial_wait = Mock()
        initial_wait.until.return_value = toolbar_button
        modal_wait = Mock()
        modal_wait.until.side_effect = lambda condition: condition(driver)
        file_list_wait = Mock()
        polls = {"count": 0}

        def wait_for_rows(condition):
            result = condition(driver)
            if not result:
                polls["count"] += 1
                result = condition(driver)
            return result

        file_list_wait.until.side_effect = wait_for_rows
        download_wait = Mock()
        download_wait.until.return_value = download_button
        driver = Mock()
        driver.find_elements.return_value = [modal]

        row_scans = {"count": 0}

        def execute_script(script, target):
            if "const matches = []" in script:
                row_scans["count"] += 1
                return [] if row_scans["count"] == 1 else photo_rows
            if target is download_button:
                create_zip()

        driver.execute_script.side_effect = execute_script

        with (
            tempfile.TemporaryDirectory() as directory,
            patch(
                "salesforce.WebDriverWait",
                side_effect=[modal_wait, file_list_wait, download_wait],
            ),
            patch("salesforce.time.monotonic", side_effect=[0, 0, 0, 3, 3]),
            patch("salesforce.time.sleep"),
        ):
            result = _baixar_arquivos_relacionados(
                driver,
                initial_wait,
                "1030177305",
                Path(directory),
            )

            self.assertEqual(result, ["1030177305_foto.jpg"])
            self.assertEqual(
                (Path(directory) / "1030177305_foto.jpg").read_bytes(),
                b"image",
            )
            self.assertFalse((Path(directory) / "download.zip").exists())

        driver.execute_script.assert_any_call(
            "arguments[0].click()",
            checkbox,
        )
        driver.execute_script.assert_any_call(
            "arguments[0].click()",
            download_button,
        )
        self.assertEqual(polls["count"], 1)
        self.assertGreaterEqual(row_scans["count"], 3)

    def test_direct_download_from_files_button_is_extracted_without_modal(self):
        toolbar_button = Mock()

        def create_zip():
            with zipfile.ZipFile(Path(directory) / "download.zip", "w") as archive:
                archive.writestr("equipe/foto.jpg", b"image")
                archive.writestr("equipe/relatorio.pdf", b"document")

        toolbar_button.click.side_effect = create_zip
        initial_wait = Mock()
        initial_wait.until.return_value = toolbar_button
        download_wait = Mock()
        download_wait.until.side_effect = lambda condition: condition(driver)
        driver = Mock()
        driver.find_elements.return_value = []

        with (
            tempfile.TemporaryDirectory() as directory,
            patch("salesforce.WebDriverWait", return_value=download_wait),
            patch("salesforce.time.monotonic", side_effect=[0, 0, 0, 3, 3]),
            patch("salesforce.time.sleep"),
        ):
            result = _baixar_arquivos_relacionados(
                driver,
                initial_wait,
                "1030177305",
                Path(directory),
            )

            self.assertEqual(result, ["1030177305_foto.jpg"])
            self.assertEqual(
                (Path(directory) / "1030177305_foto.jpg").read_bytes(),
                b"image",
            )
            self.assertFalse((Path(directory) / "download.zip").exists())

        toolbar_button.click.assert_called_once()
        driver.find_elements.assert_called_once()

    def test_empty_file_table_with_headers_returns_no_images(self):
        toolbar_button = Mock()
        close_button = Mock()
        modal = Mock()
        modal.is_displayed.return_value = True
        modal.text = "Arquivo Tipo Nenhum arquivo Baixar"
        modal.find_elements.side_effect = lambda _by, selector: (
            [Mock()]
            if selector == ".//button[normalize-space(.)='Baixar']"
            else [close_button]
            if selector == ".//button[contains(., 'Cancelar e fechar')]"
            else []
        )
        initial_wait = Mock()
        initial_wait.until.return_value = toolbar_button
        modal_wait = Mock()
        modal_wait.until.side_effect = lambda condition: condition(driver)
        file_list_wait = Mock()
        file_list_wait.until.side_effect = lambda condition: condition(driver)
        driver = Mock()
        driver.find_elements.return_value = [modal]
        driver.execute_script.return_value = []

        with (
            tempfile.TemporaryDirectory() as directory,
            patch(
                "salesforce.WebDriverWait",
                side_effect=[modal_wait, file_list_wait],
            ),
        ):
            result = _baixar_arquivos_relacionados(
                driver,
                initial_wait,
                "1030177305",
                Path(directory),
            )

        self.assertEqual(result, [])
        self.assertEqual(driver.execute_script.call_count, 3)
        driver.execute_script.assert_any_call(
            "arguments[0].click()",
            close_button,
        )

    def test_reports_failure_to_open_the_related_tab(self):
        with tempfile.TemporaryDirectory() as directory:
            driver = Mock()
            with (
                patch("salesforce.WebDriverWait", return_value=Mock()),
                patch("salesforce._buscar_ordem"),
                patch(
                    "salesforce._abrir_aba_relacionados",
                    side_effect=RuntimeError("aba Relacionados indisponível"),
                ),
                patch("salesforce._compromissos_relacionados"),
                patch("salesforce._ler_motivo", return_value=""),
                patch(
                    "salesforce._baixar_arquivos_relacionados"
                ) as download_files,
            ):
                result = _processar_nota(
                    driver,
                    Path(directory),
                    "1030177305",
                )

            self.assertEqual(result.status, "Erro")
            self.assertEqual(result.erro, "aba Relacionados indisponível")
            download_files.assert_not_called()


if __name__ == "__main__":
    unittest.main()
